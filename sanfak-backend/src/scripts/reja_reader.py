import sys, json, re
from pathlib import Path
from dataclasses import dataclass, field
from typing import Any, Optional
import openpyxl


MIN_PATTERN_RATIO = 0.6
CONFIDENCE_THRESHOLD = 0.5
MIN_SAMPLE_ROWS = 3
MAX_SAMPLE_ROWS = 10

MIN_TITLE_LENGTH = 10
MIN_TITLE_WORDS = 2


class ParseError(Exception):
    pass


class ValidationError(Exception):
    pass


def merge_xaritasi(ws):
    xarita = {}
    for rng in ws.merged_cells.ranges:
        bosh = ws.cell(rng.min_row, rng.min_col).value
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                xarita[(r, c)] = bosh
    return xarita


def v(ws, xarita, row, col):
    return xarita.get((row, col), ws.cell(row, col).value)


def qator(ws, xarita, row):
    return [v(ws, xarita, row, c) for c in range(1, ws.max_column + 1)]


def int_yoki_nol(val):
    try:
        return int(float(str(val))) if val is not None else 0
    except (TypeError, ValueError):
        return 0


def normalize(s):
    if not s:
        return ""
    return (str(s).lower()
            .replace('`',      '')
            .replace('\u2019', '')
            .replace('\u2018', '')
            .replace('\u201a', '')
            .replace('\u2011', '')
            .replace("'",      '')
            .replace('-',      '')
            .replace('\n',     '')
            .replace('\r',     '')
            .replace(' ',      '')
            .strip())


_SYMBOL_TO_WORD = {
    "%": "foiz",
    "№": "raqam",
}


def _slugify(text: str) -> str:
    if not text:
        return ""
    s = str(text).lower().strip()
    for sym, word in _SYMBOL_TO_WORD.items():
        if sym in s:
            s = s.replace(sym, f" {word} ")
    for ch in ("'", "\u2019", "\u2018", "`", "\u00B4", "\u02BC"):
        s = s.replace(ch, "")
    s = re.sub(r"[^\w\s\u0400-\u04FF]", " ", s, flags=re.UNICODE)
    s = re.sub(r"\s+", "_", s.strip())
    s = re.sub(r"_+", "_", s)
    return s.strip("_")


_SERIAL_RE = re.compile(r'^\d+(?:\.\d+)+$')
_CODE_RE   = re.compile(r'^[A-Za-zА-Яа-я][\w\-\'\u2019 /()]*\d+')
_PERCENT_RE = re.compile(r'^[\d.,]+\s*%?$')


@dataclass
class ColumnProfile:
    col: int
    header_text: str = ""
    sample_values: list[Any] = field(default_factory=list)
    n_sample: int = 0
    n_numeric: int = 0
    n_text: int = 0
    n_empty: int = 0
    n_serial_pattern: int = 0
    n_code_pattern: int = 0
    n_long_text: int = 0
    avg_length: float = 0.0
    max_value: float = 0.0
    min_value: float = 0.0
    integer_values: bool = True

    @property
    def numeric_ratio(self) -> float:
        return self.n_numeric / self.n_sample if self.n_sample else 0.0

    @property
    def serial_ratio(self) -> float:
        return self.n_serial_pattern / self.n_sample if self.n_sample else 0.0

    @property
    def code_ratio(self) -> float:
        return self.n_code_pattern / self.n_sample if self.n_sample else 0.0

    @property
    def title_ratio(self) -> float:
        return self.n_long_text / self.n_sample if self.n_sample else 0.0


def _collect_header_text(ws, xarita, col: int, tartib_qator: int) -> str:
    seen = []
    for r in range(1, tartib_qator):
        val = v(ws, xarita, r, col)
        if val is None or not str(val).strip():
            continue
        s = str(val).strip()
        if not seen or seen[-1] != s:
            seen.append(s)
    return seen[-1] if seen else ""


def _profile_column(ws, xarita, col: int, sample_rows: list[int], tartib_qator: int) -> ColumnProfile:
    prof = ColumnProfile(col=col)
    prof.header_text = _collect_header_text(ws, xarita, col, tartib_qator)

    lengths = []
    numerics = []

    for r in sample_rows:
        val = v(ws, xarita, r, col)
        prof.sample_values.append(val)
        prof.n_sample += 1

        if val is None:
            prof.n_empty += 1
            continue

        s = str(val).strip()
        if not s:
            prof.n_empty += 1
            continue

        try:
            num = float(s.replace(',', '.'))
            prof.n_numeric += 1
            numerics.append(num)
            if num != int(num):
                prof.integer_values = False
        except (ValueError, TypeError):
            prof.n_text += 1
            lengths.append(len(s))

            if _SERIAL_RE.match(s):
                prof.n_serial_pattern += 1
            if _CODE_RE.match(s):
                prof.n_code_pattern += 1
            words = [w for w in s.split() if w]
            if len(s) >= MIN_TITLE_LENGTH and len(words) >= MIN_TITLE_WORDS:
                prof.n_long_text += 1

        if _SERIAL_RE.match(s):
            pass

    prof.n_serial_pattern = sum(
        1 for vv in prof.sample_values
        if vv is not None and _SERIAL_RE.match(str(vv).strip())
    )

    if lengths:
        prof.avg_length = sum(lengths) / len(lengths)
    if numerics:
        prof.max_value = max(numerics)
        prof.min_value = min(numerics)
    return prof


def _find_sample_rows(ws, xarita, tartib_qator: int) -> list[int]:
    start = tartib_qator + 1
    end = min(start + MAX_SAMPLE_ROWS, ws.max_row + 1)
    rows = []
    for r in range(start, end):
        has_data = any(v(ws, xarita, r, c) is not None for c in range(1, min(5, ws.max_column + 1)))
        if has_data:
            rows.append(r)
        if len(rows) >= MAX_SAMPLE_ROWS:
            break
    if len(rows) < MIN_SAMPLE_ROWS:
        for r in range(start, ws.max_row + 1):
            if r not in rows:
                rows.append(r)
            if len(rows) >= MIN_SAMPLE_ROWS:
                break
    return rows


def analyze_xlsx(ws, xarita, tartib_qator: int) -> dict[int, ColumnProfile]:
    sample_rows = _find_sample_rows(ws, xarita, tartib_qator)
    return {
        col: _profile_column(ws, xarita, col, sample_rows, tartib_qator)
        for col in range(1, ws.max_column + 1)
    }


ROLE_SERIAL     = "serialNumber"
ROLE_CODE       = "code"
ROLE_TITLE      = "title"
ROLE_PARTICLE   = "particle"
ROLE_SEM_START  = "semester_start"
ROLE_SEM_END    = "semester_end"
ROLE_CRED_START = "credit_start"
ROLE_CRED_END   = "credit_end"
ROLE_TOTAL_CRED = "totalCredit"


def _detect_by_content(prof: ColumnProfile) -> tuple[Optional[str], float]:
    if prof.n_sample == 0:
        return None, 0.0

    if prof.serial_ratio >= MIN_PATTERN_RATIO:
        return ROLE_SERIAL, 0.95

    if prof.code_ratio >= MIN_PATTERN_RATIO:
        return ROLE_CODE, 0.85

    if prof.title_ratio >= MIN_PATTERN_RATIO:
        return ROLE_TITLE, 0.85

    if prof.numeric_ratio >= 0.5 and prof.n_numeric >= 3:
        return ROLE_PARTICLE, 0.6

    return None, 0.0


def _detect_by_header_sequence(profiles: dict[int, ColumnProfile],
                                ws, xarita, tartib_qator: int) -> dict[int, tuple[str, float]]:
    result = {}
    max_col = ws.max_column

    for r in range(2, tartib_qator):
        seqs = []
        c = 1
        while c <= max_col:
            val = v(ws, xarita, r, c)
            try:
                n = int(str(val).strip()) if val is not None else None
            except (ValueError, TypeError):
                n = None
            if n == 1:
                seq_start = c
                seq_vals = [1]
                cc = c + 1
                while cc <= max_col:
                    vv = v(ws, xarita, r, cc)
                    try:
                        nn = int(str(vv).strip()) if vv is not None else None
                    except (ValueError, TypeError):
                        nn = None
                    if nn == seq_vals[-1] + 1:
                        seq_vals.append(nn)
                        cc += 1
                    else:
                        break
                if len(seq_vals) >= 4:
                    seqs.append((seq_start, len(seq_vals)))
                c = cc
            else:
                c += 1

        if len(seqs) >= 2:
            s1, l1 = seqs[0]
            s2, l2 = seqs[1]
            for col in range(s1, s1 + l1):
                result[col] = (ROLE_PARTICLE if False else "semester_soat", 0.9)
            for col in range(s2, s2 + l2):
                result[col] = ("semester_kredit", 0.9)
            result[s1] = (ROLE_SEM_START, 0.95)
            result[s1 + l1 - 1] = (ROLE_SEM_END, 0.95)
            result[s2] = (ROLE_CRED_START, 0.95)
            result[s2 + l2 - 1] = (ROLE_CRED_END, 0.95)
            return result
        elif len(seqs) == 1:
            s1, l1 = seqs[0]
            result[s1] = (ROLE_SEM_START, 0.85)
            result[s1 + l1 - 1] = (ROLE_SEM_END, 0.85)
            return result

    return result


def _detect_last_numeric_column(profiles: dict[int, ColumnProfile],
                                 sem_end: Optional[int]) -> Optional[int]:
    if sem_end is None:
        candidates = sorted(profiles.keys(), reverse=True)
    else:
        candidates = sorted([c for c in profiles.keys() if c > sem_end], reverse=True)

    for col in candidates:
        prof = profiles[col]
        if prof.numeric_ratio >= 0.5 and prof.n_numeric >= 3:
            return col
    return None


def _detect_by_position(
    profiles: dict[int, ColumnProfile],
    assigned: dict[int, str],
) -> dict[int, tuple[str, float]]:
    result = {}
    serial_col = next((c for c, r in assigned.items() if r == ROLE_SERIAL), None)
    sem_start = next((c for c, r in assigned.items() if r == ROLE_SEM_START), None)

    if serial_col is None:
        return result

    cand_code = serial_col + 1
    if cand_code in profiles and cand_code not in assigned:
        prof = profiles[cand_code]
        if prof.code_ratio >= 0.3 or prof.numeric_ratio < 0.3:
            result[cand_code] = (ROLE_CODE, 0.7)

    title_start = max(serial_col, cand_code) + 1
    if title_start in profiles and title_start not in assigned:
        prof = profiles[title_start]
        if prof.title_ratio >= 0.3:
            result[title_start] = (ROLE_TITLE, 0.7)

    title_col = next(
        (c for c, r in list(assigned.items()) + list(result.items())
         if (r[0] if isinstance(r, tuple) else r) == ROLE_TITLE),
        title_start,
    )
    end = sem_start if sem_start else max(profiles.keys()) + 1
    for col in range(title_col + 1, end):
        if col in assigned or col in result:
            continue
        prof = profiles[col]
        if prof.numeric_ratio >= 0.3 or prof.n_empty >= prof.n_sample // 2:
            result[col] = (ROLE_PARTICLE, 0.6)

    return result


def detect_roles(ws, xarita, tartib_qator: int, profiles: Optional[dict] = None) -> dict:
    if profiles is None:
        profiles = analyze_xlsx(ws, xarita, tartib_qator)

    assigned = {}
    confidences = {}
    warnings = []

    content_votes: dict[str, list[tuple[int, float]]] = {}
    for col, prof in profiles.items():
        role, conf = _detect_by_content(prof)
        if role and role in (ROLE_SERIAL, ROLE_CODE, ROLE_TITLE):
            content_votes.setdefault(role, []).append((col, conf))

    for role, candidates in content_votes.items():
        candidates.sort(key=lambda x: -x[1])
        best_col, best_conf = candidates[0]
        if best_conf >= CONFIDENCE_THRESHOLD:
            assigned[best_col] = role
            confidences[role] = best_conf

    seq_results = _detect_by_header_sequence(profiles, ws, xarita, tartib_qator)
    for col, (role, conf) in seq_results.items():
        if role in (ROLE_SEM_START, ROLE_SEM_END, ROLE_CRED_START, ROLE_CRED_END):
            if role not in confidences or confidences[role] < conf:
                for c, r in list(assigned.items()):
                    if r == role:
                        del assigned[c]
                assigned[col] = role
                confidences[role] = conf

    pos_results = _detect_by_position(profiles, assigned)
    particle_cols = set()
    for col, (role, conf) in pos_results.items():
        if role == ROLE_PARTICLE:
            particle_cols.add(col)
        elif role in (ROLE_CODE, ROLE_TITLE):
            if role not in confidences:
                assigned[col] = role
                confidences[role] = conf

    sem_end = next((c for c, r in assigned.items() if r == ROLE_SEM_END), None)
    cred_end = next((c for c, r in assigned.items() if r == ROLE_CRED_END), None)
    last_numeric_after = cred_end or sem_end
    total_col = _detect_last_numeric_column(profiles, last_numeric_after)
    if total_col and total_col not in assigned:
        assigned[total_col] = ROLE_TOTAL_CRED
        confidences[ROLE_TOTAL_CRED] = 0.7

    roles = {}
    for col, role in assigned.items():
        if role in (ROLE_SERIAL, ROLE_CODE, ROLE_TITLE, ROLE_TOTAL_CRED,
                    ROLE_SEM_START, ROLE_SEM_END, ROLE_CRED_START, ROLE_CRED_END):
            roles[role] = col

    title_col = roles.get(ROLE_TITLE)
    sem_start = roles.get(ROLE_SEM_START)
    if title_col and sem_start:
        pc = [c for c in range(title_col + 1, sem_start) if c not in assigned.values()
              or assigned.get(c) == ROLE_PARTICLE]
        roles["particle_cols"] = sorted(set(pc) | particle_cols)
    else:
        roles["particle_cols"] = sorted(particle_cols)

    roles["_confidences"] = confidences
    roles["_warnings"] = warnings
    return roles


@dataclass
class ValidationResult:
    passed: bool
    warnings: list[str] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)


def validate_roles(roles: dict, profiles: dict[int, ColumnProfile]) -> ValidationResult:
    result = ValidationResult(passed=True)
    confidences = roles.get("_confidences", {})

    if "serialNumber" not in roles:
        result.errors.append("MAJBURIY: serialNumber ustuni topilmadi")
        result.passed = False
    if "title" not in roles:
        result.errors.append("MAJBURIY: title ustuni topilmadi")
        result.passed = False

    serial = roles.get("serialNumber")
    code   = roles.get("code")
    title  = roles.get("title")
    sem_s  = roles.get("semester_start")
    cred_s = roles.get("credit_start")
    total  = roles.get("totalCredit")

    if serial is not None and title is not None and serial >= title:
        result.warnings.append(
            f"Nostandart tartib: serialNumber(C{serial}) title(C{title}) dan keyin"
        )
    if code is not None and title is not None and code >= title:
        result.warnings.append(
            f"Nostandart tartib: code(C{code}) title(C{title}) dan keyin"
        )
    if title is not None and sem_s is not None and title >= sem_s:
        result.warnings.append(
            f"title(C{title}) semester_start(C{sem_s}) dan keyin"
        )
    if sem_s is not None and cred_s is not None and sem_s >= cred_s:
        result.warnings.append(
            f"semester_start(C{sem_s}) credit_start(C{cred_s}) dan keyin"
        )
    if total is not None and cred_s is not None and total <= cred_s:
        result.warnings.append(
            f"totalCredit(C{total}) credit block ichida yoki oldin"
        )

    cols_used = {}
    for role, col in roles.items():
        if role.startswith("_") or role == "particle_cols":
            continue
        if not isinstance(col, int):
            continue
        if col in cols_used:
            result.errors.append(
                f"C{col} bir nechta rolda: {cols_used[col]}, {role}"
            )
            result.passed = False
        else:
            cols_used[col] = role

    for role, conf in confidences.items():
        if conf < CONFIDENCE_THRESHOLD:
            result.warnings.append(
                f"Past confidence: {role}={conf:.2f} (< {CONFIDENCE_THRESHOLD})"
            )

    pc = roles.get("particle_cols", [])
    if not pc:
        result.warnings.append("Particle ustunlari topilmadi")
    elif len(pc) < 2:
        result.warnings.append(f"Kam particle ustun: {len(pc)}")

    return result


_DEFAULT_CANONICAL_ORDER = [
    "hour",
    "percent",
    "total",
    "lecture",
    "practical",
    "laboratory",
    "seminar",
    "courseWork",
    "independent",
]

_CANONICAL_BY_HEADER = [
    (re.compile(r"klinik", re.I), "clinicalPractice"),
    (re.compile(r"kurs\s*ishi", re.I), "courseWork"),
    (re.compile(r"mustaqil", re.I), "independent"),
]


def _canonical_by_header(header: Optional[str]) -> Optional[str]:
    if not header:
        return None
    for rx, canonical in _CANONICAL_BY_HEADER:
        if rx.search(str(header)):
            return canonical
    return None


def _infer_canonical_by_content(
    particle_cols: list[int],
    profiles: dict[int, ColumnProfile],
) -> dict[int, str]:
    if not particle_cols:
        return {}

    result = {}
    sorted_cols = sorted(particle_cols)

    if len(sorted_cols) == 9:
        for i, col in enumerate(sorted_cols):
            result[col] = _DEFAULT_CANONICAL_ORDER[i]
        return result

    if len(sorted_cols) == 8:
        order = ["hour", "percent", "total", "lecture", "practical",
                 "laboratory", "seminar", "independent"]
        for i, col in enumerate(sorted_cols):
            result[col] = order[i]
        return result

    for i, col in enumerate(sorted_cols):
        result[col] = _DEFAULT_CANONICAL_ORDER[i] if i < len(_DEFAULT_CANONICAL_ORDER) else None

    return result


def build_particle_items(
    roles: dict,
    profiles: dict[int, ColumnProfile],
    ws=None,
    xarita=None,
    tartib_qator: Optional[int] = None,
) -> list[dict]:
    particle_cols = sorted(roles.get("particle_cols", []))
    canonical_map = _infer_canonical_by_content(particle_cols, profiles)

    items = []
    for col in particle_cols:
        prof = profiles.get(col)
        if prof is None:
            continue
        header = prof.header_text or f"col_{col}"
        canonical = _canonical_by_header(prof.header_text) or canonical_map.get(col)
        slug = _slugify(header) or canonical or f"col_{col}"
        col_num = None
        if ws is not None and xarita is not None and tartib_qator:
            raw = v(ws, xarita, tartib_qator, col)
            try:
                col_num = int(str(raw).strip()) if raw is not None else None
            except (TypeError, ValueError):
                col_num = None
        items.append({
            "slug":      slug,
            "canonical": canonical,
            "title":     header,
            "col":       col,
            "colNum":    col_num,
        })
    return items


def tartib_qatorini_top(ws, xarita):
    for r in range(1, min(25, ws.max_row)):
        vals = qator(ws, xarita, r)
        raqamlar = []
        for val in vals:
            try:
                raqamlar.append(int(str(val).strip()))
            except (TypeError, ValueError):
                pass
        if len(raqamlar) >= 5 and sorted(raqamlar)[:5] == list(range(1, 6)):
            return r
    return None


def _detect_kurs_names(ws, xarita, soat_boshliq: int, sem_soni: int, tartib_qator: int) -> list:
    kurs_soni = max(1, sem_soni // 2)
    names = []

    for r in range(2, tartib_qator):
        row_names = []
        for k in range(kurs_soni):
            c = soat_boshliq + k * 2
            val = v(ws, xarita, r, c)
            if not val or not str(val).strip():
                continue
            s = str(val).strip()
            next_val = v(ws, xarita, r, c + 1) if c + 1 <= ws.max_column else None
            if str(next_val or "").strip() == s:
                row_names.append(s)
        if len(row_names) == kurs_soni and len(set(row_names)) == kurs_soni:
            return row_names
        if row_names and len(set(row_names)) > 1 and not names:
            names = row_names
    return names


def sarlavha_meta_qur(ws, xarita, tartib_qator, ust=None, roles=None, particle_items=None):

    def col_val(col):
        if not col:
            return None
        oxirgi_ustun = ws.max_column
        for r in range(2, tartib_qator):
            val = v(ws, xarita, r, col)
            if val and str(val).strip():
                s = str(val).strip()
                birinchi = v(ws, xarita, r, 1)
                oxirgi = v(ws, xarita, r, oxirgi_ustun)
                if birinchi == val and oxirgi == val and oxirgi_ustun > col:
                    continue
                return s
        return None

    def col_val_specific(col):
        if not col:
            return None
        last_val = None
        for r in range(tartib_qator - 1, 1, -1):
            val = v(ws, xarita, r, col)
            if val and str(val).strip():
                s = str(val).strip()
                left_val = v(ws, xarita, r, col - 1) if col > 1 else None
                right_val = v(ws, xarita, r, col + 1)
                if str(left_val or "").strip() == s or str(right_val or "").strip() == s:
                    continue
                return s
        return None

    def range_val(c1, c2):
        if not c1 or not c2:
            return None
        oxirgi_ustun = ws.max_column
        for r in range(2, tartib_qator):
            val = v(ws, xarita, r, c1)
            if val and str(val).strip():
                s = str(val).strip()
                if not all(v(ws, xarita, r, c) == val for c in range(c1, c2 + 1)):
                    continue
                chap = v(ws, xarita, r, c1 - 1) if c1 > 1 else None
                ong = v(ws, xarita, r, c2 + 1) if c2 < oxirgi_ustun else None
                if chap == val or ong == val:
                    continue
                return s
        return None

    max_col = ws.max_column
    ust = ust or {}

    ustunlar = {}
    for col in range(1, min(12, max_col + 1)):
        nom = col_val(col)
        if nom:
            ustunlar[col] = nom

    soat_boshliq = kredit_boshliq = None
    sem_soni_detected = 12
    for r in range(2, tartib_qator):
        vals = qator(ws, xarita, r)
        col_count = len(vals)
        for start in range(col_count):
            seq = []
            for i in range(start, col_count):
                b = vals[i]
                try:
                    n = int(b)
                    if seq and n != seq[-1] + 1:
                        break
                    seq.append(n)
                except (TypeError, ValueError):
                    break
            if len(seq) >= 4 and seq[0] == 1:
                if soat_boshliq is None:
                    soat_boshliq      = start + 1
                    sem_soni_detected = len(seq)
                elif kredit_boshliq is None and start + 1 != soat_boshliq:
                    kredit_boshliq = start + 1
                    break
        if soat_boshliq and kredit_boshliq:
            break

    meta_sarlavha = {
        "serialNumber" : col_val(ust.get("serialNumber")) or ustunlar.get(1, ""),
        "code"         : col_val(ust.get("code"))         or ustunlar.get(2, ""),
        "title"        : col_val(ust.get("title"))        or ustunlar.get(3, ""),
    }

    if particle_items is None:
        particle_items = []
    if not particle_items and ust:
        for canonical in ("hour", "percent", "total", "lecture", "practical",
                          "laboratory", "seminar", "courseWork", "independent"):
            col_idx = ust.get(canonical)
            if col_idx:
                title = col_val_specific(col_idx) or canonical
                particle_items.append({
                    "slug":      _slugify(title) or canonical,
                    "title":     title,
                    "canonical": canonical,
                    "col":       col_idx,
                })

    p_cols = [it["col"] for it in particle_items if it.get("col")]
    yuk_c1 = min(p_cols, default=4)
    yuk_c2 = max(p_cols, default=11)
    yuk_nom = range_val(yuk_c1, yuk_c2)

    meta_sarlavha["particles"] = {
        "title": yuk_nom or "",
        "items": [
            {
                "slug": it["slug"],
                "title": it["title"],
                **({"colNum": it["colNum"]} if it.get("colNum") is not None else {}),
            }
            for it in particle_items
        ],
    }

    if soat_boshliq:
        soat_blok_nom   = range_val(soat_boshliq, soat_boshliq + sem_soni_detected - 1)
        kredit_blok_nom = (range_val(kredit_boshliq, kredit_boshliq + sem_soni_detected - 1)
                           if kredit_boshliq else None)

        kurs_soni = max(1, sem_soni_detected // 2)
        kurs_nomlari_s = _detect_kurs_names(ws, xarita, soat_boshliq, sem_soni_detected, tartib_qator)
        kurs_nomlari_k = _detect_kurs_names(ws, xarita, kredit_boshliq, sem_soni_detected, tartib_qator) if kredit_boshliq else []

        hafta_soat   = _hafta_sonlarini_o_qi(ws, xarita, tartib_qator, soat_boshliq, sem_soni_detected,
                                              label_keywords=[_HAFTA_SEMESTR_LABEL])
        hafta_kred   = (_hafta_sonlarini_o_qi(ws, xarita, tartib_qator, kredit_boshliq, sem_soni_detected,
                                               label_keywords=[_KREDIT_TAQSIMOT_LABEL])
                        if kredit_boshliq else [30] * sem_soni_detected)

        hafta_kurs_qatori_s = _hafta_sonlarini_o_qi(ws, xarita, tartib_qator, soat_boshliq, sem_soni_detected,
                                                     label_keywords=[_HAFTA_KURS_LABEL])
        hafta_kurs_qatori_k = (_hafta_sonlarini_o_qi(ws, xarita, tartib_qator, kredit_boshliq, sem_soni_detected,
                                                      label_keywords=[_HAFTA_KURS_LABEL])
                                if kredit_boshliq else [])
        hafta_kurs_s = ([hafta_kurs_qatori_s[i] for i in range(0, len(hafta_kurs_qatori_s), 2)]
                         if hafta_kurs_qatori_s else [hafta_soat[i] for i in range(0, len(hafta_soat), 2)])
        hafta_kurs_k = ([hafta_kurs_qatori_k[i] for i in range(0, len(hafta_kurs_qatori_k), 2)]
                         if hafta_kurs_qatori_k else [hafta_kred[i] for i in range(0, len(hafta_kred), 2)])

        meta_sarlavha["distribution"] = {
            "title":    soat_blok_nom or "",
            "courses":  kurs_nomlari_s or [str(i + 1) for i in range(kurs_soni)],
            "weekly":   hafta_kurs_s or [30] * kurs_soni,
            "semester": list(range(1, sem_soni_detected + 1)),
            "audience": hafta_soat or [15] * sem_soni_detected,
        }

        meta_sarlavha["credit"] = {
            "title":    kredit_blok_nom or "",
            "courses":  kurs_nomlari_s or [str(i + 1) for i in range(kurs_soni)],
            "weekly":   hafta_kurs_k or [30] * kurs_soni,
            "semester": list(range(1, sem_soni_detected + 1)),
            "distribution": hafta_kred or [30] * sem_soni_detected,
        }

        def _col_num(col):
            if not col:
                return None
            raw = v(ws, xarita, tartib_qator, col)
            try:
                return int(str(raw).strip()) if raw is not None else None
            except (TypeError, ValueError):
                return None

        meta_sarlavha["distribution"]["semesterColNums"] = [
            _col_num(soat_boshliq + i) for i in range(sem_soni_detected)
        ]
        meta_sarlavha["distribution"]["courseColNums"] = [
            _col_num(soat_boshliq + k * 2) for k in range(kurs_soni)
        ]
        if kredit_boshliq:
            meta_sarlavha["credit"]["semesterColNums"] = [
                _col_num(kredit_boshliq + i) for i in range(sem_soni_detected)
            ]
            meta_sarlavha["credit"]["courseColNums"] = [
                _col_num(kredit_boshliq + k * 2) for k in range(kurs_soni)
            ]

    meta_sarlavha["totalCredit"] = col_val(max_col) or ""

    def _col_num2(col):
        if not col:
            return None
        raw = v(ws, xarita, tartib_qator, col)
        try:
            return int(str(raw).strip()) if raw is not None else None
        except (TypeError, ValueError):
            return None

    columns = {}
    if ust and ust.get("serialNumber"):
        columns["serialNumber"] = _col_num2(ust["serialNumber"])
    if ust and ust.get("code"):
        columns["code"] = _col_num2(ust["code"])
    if ust and ust.get("title"):
        columns["title"] = _col_num2(ust["title"])
    if max_col:
        columns["totalCredit"] = _col_num2(max_col)
    if soat_boshliq:
        columns["distributionStart"] = _col_num2(soat_boshliq)
        columns["distributionEnd"]   = _col_num2(soat_boshliq + sem_soni_detected - 1)
    if kredit_boshliq:
        columns["creditStart"] = _col_num2(kredit_boshliq)
        columns["creditEnd"]   = _col_num2(kredit_boshliq + sem_soni_detected - 1)
    columns = {k: v for k, v in columns.items() if v is not None}
    if columns:
        meta_sarlavha["columns"] = columns

    tartib_row = []
    if tartib_qator:
        for c in range(1, ws.max_column + 1):
            raw = v(ws, xarita, tartib_qator, c)
            try:
                n = int(str(raw).strip()) if raw is not None else None
            except (TypeError, ValueError):
                n = None
            if n is not None:
                tartib_row.append(n)
    if tartib_row:
        meta_sarlavha["tartibRow"] = tartib_row

    return meta_sarlavha, soat_boshliq, kredit_boshliq, sem_soni_detected, particle_items


_HAFTA_KURS_LABEL = normalize("Kurslardagi haftalar soni")
_HAFTA_SEMESTR_LABEL = normalize("Semestrdagi auditoriya mashg'ulotlari haftalarining soni")
_KREDIT_TAQSIMOT_LABEL = normalize("Kredit taqsimoti")


def _hafta_sonlarini_o_qi(ws, xarita, tartib_qator, boshliq_col, sem_soni, label_keywords=None):
    if not boshliq_col:
        return []

    def _vals_at(r):
        vals = []
        for i in range(sem_soni):
            val = v(ws, xarita, r, boshliq_col + i)
            try:
                vals.append(int(val))
            except (TypeError, ValueError):
                vals.append(None)
        return vals

    def _accept(vals):
        numeric = [x for x in vals if x is not None]
        if len(numeric) >= sem_soni // 2 and all(1 <= x <= 100 for x in numeric):
            avg = round(sum(numeric) / len(numeric))
            return [x if x is not None else avg for x in vals]
        return None

    if label_keywords:
        for r in range(2, tartib_qator):
            label = normalize(v(ws, xarita, r, boshliq_col))
            if label and any(kw in label for kw in label_keywords):
                result = _accept(_vals_at(r + 1))
                if result is not None:
                    return result

    for r in range(2, tartib_qator):
        result = _accept(_vals_at(r))
        if result is not None:
            return result
    return []


def ustun_indekslar(ws, xarita, tartib_qator):
    profiles = analyze_xlsx(ws, xarita, tartib_qator)
    roles = detect_roles(ws, xarita, tartib_qator, profiles)
    items = build_particle_items(roles, profiles, ws, xarita, tartib_qator)
    return _roles_to_ust(roles, items)


def fan_obyekti(ws, xarita, row, ust, soat_boshliq, kredit_boshliq, sem_soni, particle_items=None):
    def get(kalit):
        col = ust.get(kalit)
        return v(ws, xarita, row, col) if col else None

    semesters = {}
    for i in range(sem_soni):
        hour_val   = int_yoki_nol(v(ws, xarita, row, soat_boshliq   + i)) if soat_boshliq   else 0
        credit_val = int_yoki_nol(v(ws, xarita, row, kredit_boshliq + i)) if kredit_boshliq else 0
        semesters[str(i + 1)] = {
            "hour"   : hour_val,
            "credit" : credit_val,
        }

    particle = []
    for it in (particle_items or []):
        col_idx = it.get("col")
        if not col_idx and ust:
            kalit = it.get("canonical") or it.get("kalit") or it["slug"]
            col_idx = ust.get(kalit)
        value = int_yoki_nol(v(ws, xarita, row, col_idx)) if col_idx else 0
        entry = {
            "slug":  it["slug"],
            "title": it.get("title", ""),
            "value": value,
        }
        if it.get("canonical"):
            entry["canonical"] = it["canonical"]
        if it.get("colNum") is not None:
            entry["colNum"] = it["colNum"]
        particle.append(entry)

    return {
        "serialNumber" : str(get("serialNumber") or "").strip(),
        "code"         : str(get("code")         or "").strip(),
        "title"        : str(get("title")        or "").strip(),
        "particle"     : particle,
        "semesters"    : semesters,
        "totalCredit"  : int_yoki_nol(get("totalCredit")),
    }


def blok_tuzilmasi_qur(ws, xarita, data_boshliq, ust, soat_boshliq, kredit_boshliq, sem_soni, particle_items=None):
    bloklar    = {}
    joriy_blok = None
    totals     = []

    AGGREGATE_TITLES = {"jami", "hammasi", "jamisemestrda", "jamikreditlar"}
    PRACTICE_SECTIONS = {"malakaviyamaliyot", "amaliyot", "malakaviyamaliyotlar"}

    for row in range(data_boshliq, ws.max_row + 1):
        t_r = str(v(ws, xarita, row, ust.get("serialNumber", 1)) or "").strip()
        kod = str(v(ws, xarita, row, ust.get("code",         2)) or "").strip()
        nom = str(v(ws, xarita, row, ust.get("title",        3)) or "").strip()

        if not nom and not kod:
            continue

        fan = fan_obyekti(ws, xarita, row, ust, soat_boshliq, kredit_boshliq, sem_soni, particle_items)

        try:
            f_tr    = float(t_r)
            is_blok = (
                abs(f_tr - round(f_tr)) < 1e-6
                and round(f_tr) >= 1
                and bool(nom or kod)
            )
        except ValueError:
            is_blok = False

        nom_norm = normalize(nom)

        if not t_r and not kod and nom_norm in AGGREGATE_TITLES:
            totals.append(fan)
            continue

        if not t_r and nom_norm in PRACTICE_SECTIONS:
            block_id = kod or "MA"
            joriy_blok = {**fan, "blockCode": block_id, "isPractice": True, "sciences": []}
            bloklar[block_id] = joriy_blok
            continue

        if is_blok:
            block_id     = kod if kod else f"BLK{int(round(f_tr))}"
            joriy_blok   = {**fan, "blockCode": block_id, "sciences": []}
            bloklar[block_id] = joriy_blok
        else:
            if joriy_blok is not None:
                joriy_blok["sciences"].append(fan)
            else:
                bloklar[kod or nom] = {**fan, "sciences": []}

    return bloklar, totals


def _roles_to_ust(roles: dict, particle_items: list[dict]) -> dict:
    ust = {}
    for key in ("serialNumber", "code", "title", "totalCredit"):
        if key in roles:
            ust[key] = roles[key]

    for it in particle_items:
        canonical = it.get("canonical")
        if canonical:
            ust[canonical] = it["col"]

    return ust


IZOH_SHEET_NOMLARI = {"izoh", "izohlar", "eslatma", "eslatmalar"}


def izoh_o_qi(wb):
    nom = next((s for s in wb.sheetnames if normalize(s) in IZOH_SHEET_NOMLARI), None)
    if nom is None:
        return None

    ws = wb[nom]
    satrlar = []
    for row in ws.iter_rows(values_only=True):
        hujayralar = [str(c).strip() for c in row if c is not None and str(c).strip()]
        if not hujayralar:
            continue
        matn = " ".join(hujayralar)
        if normalize(matn.rstrip(":")) in IZOH_SHEET_NOMLARI:
            continue
        satrlar.append(matn)

    return "\n".join(satrlar) or None


def o_qi(fayl_yoli, sheet_nom=None):
    wb = openpyxl.load_workbook(fayl_yoli, data_only=True)
    ws = wb[sheet_nom] if sheet_nom else wb.active
    xarita = merge_xaritasi(ws)

    tartib_qator = tartib_qatorini_top(ws, xarita)
    if not tartib_qator:
        raise ParseError(f"Tartib raqamlari qatori topilmadi: {fayl_yoli}")

    profiles = analyze_xlsx(ws, xarita, tartib_qator)
    roles    = detect_roles(ws, xarita, tartib_qator, profiles)

    vr = validate_roles(roles, profiles)
    for w in vr.warnings:
        print(f"WARN: {w}", file=sys.stderr)
    if not vr.passed:
        for e in vr.errors:
            print(f"ERROR: {e}", file=sys.stderr)
        raise ParseError(f"Validatsiya: {vr.errors}")

    particle_items = build_particle_items(roles, profiles, ws, xarita, tartib_qator)

    ust = _roles_to_ust(roles, particle_items)

    soat_boshliq   = roles.get("semester_start")
    kredit_boshliq = roles.get("credit_start")
    sem_end        = roles.get("semester_end")
    sem_soni       = (sem_end - soat_boshliq + 1) if (sem_end and soat_boshliq) else 10

    data_boshliq = tartib_qator + 1

    meta_sarlavha, _, _, _, _ = sarlavha_meta_qur(
        ws, xarita, tartib_qator,
        ust=ust, roles=roles, particle_items=particle_items,
    )

    bloklar, totals = blok_tuzilmasi_qur(
        ws, xarita, data_boshliq, ust,
        soat_boshliq, kredit_boshliq, sem_soni, particle_items
    )
    if totals:
        meta_sarlavha = {**meta_sarlavha, "totals": totals}

    izoh = izoh_o_qi(wb)
    if izoh:
        meta_sarlavha = {**meta_sarlavha, "izoh": izoh}

    natija = {"meta": meta_sarlavha, **bloklar}

    texnik_meta = {
        "fayl"           : str(fayl_yoli),
        "sheet"          : ws.title,
        "jami_qator"     : ws.max_row,
        "jami_ustun"     : ws.max_column,
        "tartib_qator"   : tartib_qator,
        "data_boshliq"   : data_boshliq,
        "soat_boshliq"   : soat_boshliq,
        "kredit_boshliq" : kredit_boshliq,
        "sem_soni"       : sem_soni,
        "ustunlar"       : ust,
        "roles"          : {k: v for k, v in roles.items() if not k.startswith("_")},
        "warnings"       : vr.warnings,
    }

    wb.close()
    return natija, texnik_meta


def json_saqlash(natija, chiqish_fayl):
    with open(chiqish_fayl, "w", encoding="utf-8") as f:
        json.dump(natija, f, ensure_ascii=False, indent=2)
    return chiqish_fayl


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser(description="Reja Reader v4 — O'quv reja xlsx o'quvchi")
    parser.add_argument("fayl",     help="xlsx fayl yo'li")
    parser.add_argument("--sheet",  help="Sheet nomi (ixtiyoriy)", default=None)
    parser.add_argument("--stdout", help="Natijani stdout ga chiqarish — Node.js subprocess uchun",
                        action="store_true")
    parser.add_argument("--out",    help="Chiqish JSON fayl nomi (--stdout bo'lmasa)", default=None)
    args = parser.parse_args()

    natija, meta = o_qi(args.fayl, sheet_nom=args.sheet)

    if args.stdout:
        sys.stdout.write(json.dumps(natija, ensure_ascii=False))
        sys.stdout.flush()
    else:
        json_out = args.out or (Path(args.fayl).stem + "_natija.json")
        json_saqlash(natija, json_out)

        blok_soni = sum(1 for k in natija if k != "meta")
        jami_fan = sum(
            len(b.get("sciences", []))
            for k, b in natija.items()
            if k != "meta" and isinstance(b, dict)
        )

        print(f"Saqlandi     : {json_out}")
        print(f"   Sheet        : {meta['sheet']}")
        print(f"   Semestrlar   : {natija['meta']['distribution']['semester']}")
        print(f"   Kurslar      : {natija['meta']['distribution']['courses']}")
        print(f"   Bloklar      : {blok_soni}  {[k for k in natija if k != 'meta']}")
        print(f"   Fanlar       : {jami_fan}")
        print(f"   Ustunlar     : {meta['ustunlar']}")
        print(f"   Fayl hajmi   : {Path(json_out).stat().st_size / 1024:.1f} KB")

        for blok_kod, blok in natija.items():
            if blok_kod == "meta":
                continue
            if blok.get("sciences"):
                import pprint
                print(f"\nNamuna ({blok_kod} -> birinchi fan):")
                pprint.pprint(blok["sciences"][0], width=70)
                break
