import sys
import json
import re
from pathlib import Path
import openpyxl


RIM_ARAB = {
    "I": 1, "II": 2, "III": 3, "IV": 4,
    "V": 5, "VI": 6, "VII": 7, "VIII": 8, "IX": 9, "X": 10,
}

TOTAL_WORDS = ("jami", "total", "всего", "итого", "ҳаммаси", "жами", "sum")
ALL_WORDS = ("hammasi", "all", "всего", "ҳаммаси")
COURSE_WORDS = ("kurs", "course", "year", "курс", "yil")

CANONICAL_LEGEND = {
    " ": "Nazariy va amaliy ta'lim",
    "A": "Attestatsiyalar",
    "K": "Kredit ta'lim tizimiga kirish",
    "M": "Malakaviy amaliyot",
    "D": "Yakuniy Davlat attestatsiyasi",
    "T": "Ta'til",
    "G": "GPA ko'rsatkichini hisoblash",
}

HOMOGLYPH = {
    "А": "A",
    "В": "B",
    "С": "C",
    "Е": "E",
    "Н": "H",
    "К": "K",
    "М": "M",
    "О": "O",
    "Р": "P",
    "Т": "T",
    "Х": "X",
    "І": "I",
    "Ѕ": "S",
}


def _harf_normallashtir(ch: str) -> str:
    return HOMOGLYPH.get(ch, ch)


def _normallashtir(matn: str) -> str:
    matn = matn.lower()
    for ch in ("\u2019", "\u2018", "\u02BC", "\u0060", "\u00B4"):
        matn = matn.replace(ch, "'")
    while "  " in matn:
        matn = matn.replace("  ", " ")
    return matn.strip()


def _slugify(text):
    if not text:
        return "unknown"
    s = text.lower().strip()
    for ch in ("'", "\u2019", "\u2018", "`", "\u00B4", "\u02BC"):
        s = s.replace(ch, "")
    s = re.sub(r"[^\w\s]", " ", s, flags=re.UNICODE)
    s = re.sub(r"\s+", "_", s.strip("_ "))
    s = re.sub(r"_+", "_", s)
    return s or "unknown"


def _tokenize_for_match(text):
    if not text:
        return []
    s = text.lower()
    for ch in "'ʼ`\u2018\u2019":
        s = s.replace(ch, "")
    s = re.sub(r"[^\w\s]", " ", s, flags=re.UNICODE)
    return [t for t in s.split() if len(t) >= 3]


def _tokens_match(a, b, min_prefix=4):
    if a == b:
        return True
    shorter, longer = (a, b) if len(a) < len(b) else (b, a)
    return len(shorter) >= min_prefix and longer.startswith(shorter)


def _match_legend_key(header_text, legend_titles, threshold=0.5):
    if not header_text or not legend_titles:
        return None
    header_toks = _tokenize_for_match(header_text)
    if not header_toks:
        return None
    best_key = None
    best_score = 0.0
    for key, title in legend_titles.items():
        if not title or not isinstance(title, str):
            continue
        title_toks = _tokenize_for_match(title)
        if not title_toks:
            continue
        matches = sum(
            1 for tt in title_toks if any(_tokens_match(tt, ht) for ht in header_toks)
        )
        score = matches / len(title_toks)
        if score > best_score:
            best_score = score
            best_key = key
    return best_key if best_score >= threshold else None


def _is_total_word(text_lower):
    t = text_lower.strip()
    return any(t == kw or t.startswith(kw + " ") or t.startswith(kw + ":") for kw in TOTAL_WORDS)


class JarayonReader:

    def __init__(self, fayl_yoli, sheet_nom=None):
        self.fayl_yoli = fayl_yoli
        self.sheet_nom = sheet_nom
        self.wb = openpyxl.load_workbook(fayl_yoli, data_only=True)
        self.ws = self._sheet_top()
        self.xarita = self._merge_xaritasi()
        self.layout = self._discover_layout()
        self.warnings = []

    def _sheet_top(self):
        if self.sheet_nom and self.sheet_nom in self.wb.sheetnames:
            return self.wb[self.sheet_nom]
        SHEET_KEYWORDS = (
            "jarayon", "grafik", "kalendar",
            "schedule", "calendar", "process",
            "процесс", "календар", "график",
        )
        for nom in self.wb.sheetnames:
            low = nom.lower()
            if any(kw in low for kw in SHEET_KEYWORDS):
                return self.wb[nom]
        return self.wb.active

    def _merge_xaritasi(self):
        xarita = {}
        for rng in self.ws.merged_cells.ranges:
            bosh = self.ws.cell(rng.min_row, rng.min_col).value
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    xarita[(r, c)] = bosh
        return xarita

    def _v(self, r, c):
        return self.xarita.get((r, c), self.ws.cell(r, c).value)

    def _xom(self, r, c):
        return self.ws.cell(r, c).value

    def _num_val(self, val, default=0):
        if val is None:
            return default
        if isinstance(val, (int, float)):
            return int(val) if val == int(val) else val
        try:
            f = float(str(val).replace(",", "."))
            return int(f) if f == int(f) else f
        except (TypeError, ValueError):
            return default

    def _katak_matn(self, r, c):
        val = self._v(r, c)
        return _normallashtir(str(val)) if val is not None else ""

    def _discover_layout(self):
        ws    = self.ws
        max_r = ws.max_row
        max_c = ws.max_column

        layout = {
            "kurs_satri"     : None,
            "kurs_ustun"     : 1,
            "hafta_satri"    : None,
            "oy_satri"       : None,
            "hafta_xaritasi" : {},
            "oy_chegaralari" : [],
            "stat_bosh_ust"  : None,
            "stat_ustunlar"  : {},
            "kurs_qatorlari" : [],
        }

        for r in range(1, max_r + 1):
            for c in range(1, max_c + 1):
                val = self._xom(r, c)
                try:
                    if int(str(val)) != 1:
                        continue
                except (TypeError, ValueError):
                    continue
                ketma_ket = 0
                prev      = 0
                for cc in range(c, max_c + 1):
                    v = self._xom(r, cc)
                    try:
                        n = int(str(v))
                        if n == prev + 1:
                            ketma_ket += 1
                            prev = n
                        else:
                            break
                    except (TypeError, ValueError):
                        break
                min_kk = max(2, int(max_c * 0.05))
                if ketma_ket >= min_kk:
                    layout["hafta_satri"] = r
                    layout["oy_satri"]    = r - 1
                    for hc in range(c, max_c + 1):
                        hval = self._xom(r, hc)
                        try:
                            h = int(str(hval))
                            if h > 0:
                                layout["hafta_xaritasi"][h] = hc
                        except (TypeError, ValueError):
                            pass
                    break
            if layout["hafta_xaritasi"]:
                break

        if layout["hafta_satri"]:
            layout["kurs_satri"] = layout["hafta_satri"] + 1
            for r in range(layout["kurs_satri"], min(layout["kurs_satri"] + 10, max_r + 1)):
                for c in range(1, min(max_c + 1, 5)):
                    val = str(self._v(r, c) or "").strip()
                    if val in RIM_ARAB:
                        layout["kurs_ustun"] = c
                        break
                if layout["kurs_ustun"]:
                    break

        if layout["hafta_xaritasi"]:
            layout["stat_bosh_ust"] = max(layout["hafta_xaritasi"].values()) + 1

        if layout["kurs_satri"]:
            kurs_ust = layout["kurs_ustun"]
            for r in range(layout["kurs_satri"], max_r + 1):
                val = str(self._v(r, kurs_ust) or "").strip()
                if val in RIM_ARAB:
                    layout["kurs_qatorlari"].append(r)

        if layout["kurs_satri"] and layout["stat_bosh_ust"]:
            sarlavha_satrlari = list(range(1, layout["kurs_satri"]))
            stat_bosh = layout["stat_bosh_ust"]
            for c in range(stat_bosh, max_c + 1):
                seen = []
                for r in sarlavha_satrlari:
                    qiymat = self._v(r, c)
                    if qiymat and str(qiymat).strip():
                        s = str(qiymat).strip()
                        if not seen or seen[-1] != s:
                            seen.append(s)
                if not seen:
                    continue
                header = seen[-1]
                slug = _slugify(header)
                low = _normallashtir(header)
                if any(kw == low or kw in low.split() for kw in ALL_WORDS):
                    kalit = "all"
                elif _is_total_word(low):
                    kalit = "total"
                else:
                    kalit = slug
                layout["stat_ustunlar"][c] = {
                    "key": kalit,
                    "header": header,
                    "slug": slug,
                }

        oy_satri = layout.get("oy_satri")
        if oy_satri and layout["hafta_xaritasi"]:
            oxirgi_hafta_ust = max(layout["hafta_xaritasi"].values())
            kurs_ust = layout.get("kurs_ustun", 1)
            joriy_oy = None
            joriy_bosh = None
            for c in range(kurs_ust + 1, oxirgi_hafta_ust + 1):
                raw = self._v(oy_satri, c)
                if raw is None or not str(raw).strip():
                    continue
                nom = str(raw).strip()
                if _is_total_word(_normallashtir(nom)):
                    break
                if joriy_oy is None:
                    joriy_oy = nom
                    joriy_bosh = c
                elif nom != joriy_oy:
                    layout["oy_chegaralari"].append((joriy_oy, joriy_bosh, c - 1))
                    joriy_oy = nom
                    joriy_bosh = c
            if joriy_oy and joriy_bosh:
                layout["oy_chegaralari"].append((joriy_oy, joriy_bosh, oxirgi_hafta_ust))

        return layout

    def _legend_o_qi(self):
        ws = self.ws
        max_r = ws.max_row
        max_c = ws.max_column

        kurs_qatorlari = self.layout.get("kurs_qatorlari", [])
        start_r = (max(kurs_qatorlari) + 1) if kurs_qatorlari else max(1, max_r - 10)

        best_r = None
        best_count = 0
        for r in range(start_r, max_r + 1):
            count = 0
            for c in range(1, max_c + 1):
                raw = self._xom(r, c)
                if raw is None:
                    continue
                s = str(raw).strip()
                if len(s) == 1 and s.isalpha():
                    count += 1
            if count > best_count:
                best_count = count
                best_r = r

        keys = []
        if not best_r or best_count < 1:
            return self._legend_zaxira()

        blank_title = ""
        for c in range(1, max_c + 1):
            raw = self._xom(best_r, c)
            if raw is None:
                continue
            s = str(raw).strip()
            if len(s) > 3 and not (len(s) == 1 and s.isalpha()):
                blank_title = s
                break

        keys.append({"key": " ", "title": blank_title})

        for c in range(1, max_c + 1):
            raw = self._xom(best_r, c)
            if raw is None:
                continue
            s = str(raw).strip()
            if len(s) == 1 and s.isalpha():
                key = s.upper()
                title = key
                for cn in range(c + 1, max_c + 1):
                    next_raw = self._v(best_r, cn)
                    if next_raw is None:
                        continue
                    ns = str(next_raw).strip()
                    if ns and not (len(ns) == 1 and ns.isalpha()):
                        title = ns
                        break
                    if len(ns) == 1 and ns.isalpha():
                        break
                if not any(el["key"] == key for el in keys):
                    keys.append({"key": key, "title": title})

        return keys

    def _amalda_uchragan_harflar(self):
        layout = self.layout
        hafta_ustunlari = set(layout.get("hafta_xaritasi", {}).values())
        harflar = set()
        for r in layout.get("kurs_qatorlari", []):
            for c in hafta_ustunlari:
                val = self._xom(r, c)
                if val is None:
                    continue
                s = str(val).strip()
                if len(s) == 1 and s.isalpha():
                    harflar.add(_harf_normallashtir(s.upper()))
        return harflar

    def _legend_zaxira(self):
        harflar = self._amalda_uchragan_harflar()
        keys = [{"key": " ", "title": CANONICAL_LEGEND.get(" ", "")}]
        nomalum = []
        for h in sorted(harflar):
            if h in CANONICAL_LEGEND:
                keys.append({"key": h, "title": CANONICAL_LEGEND[h]})
            else:
                keys.append({"key": h, "title": h})
                nomalum.append(h)

        message = "Legend qatori faylda topilmadi — kanonik zaxira jadval ishlatildi."
        if nomalum:
            message += f" Kanonik jadvalda yo'q harflar: {', '.join(nomalum)}."
        self.warnings.append({"code": "LEGEND_FALLBACK", "message": message})
        return keys

    def _kurslar_o_qi(self):
        ws      = self.ws
        layout  = self.layout
        courses = []
        if not layout["kurs_satri"]:
            return courses

        kurs_ust = layout.get("kurs_ustun", 1)

        for r in range(layout["kurs_satri"], ws.max_row + 1):
            kurs_val = str(self._v(r, kurs_ust) or "").strip()
            if kurs_val not in RIM_ARAB:
                continue

            weeks_xom = {}
            for h_raqam, h_ustun in layout["hafta_xaritasi"].items():
                val = self._xom(r, h_ustun)
                weeks_xom[h_raqam] = str(val).strip() if val else None

            months = []
            for oy_nom, c1, c2 in layout["oy_chegaralari"]:
                weeks_list = []
                for h_raqam, h_ustun in layout["hafta_xaritasi"].items():
                    if c1 <= h_ustun <= c2:
                        key = weeks_xom.get(h_raqam)
                        weeks_list.append({
                            "week": h_raqam,
                            "key" : key if key else " ",
                        })
                weeks_list.sort(key=lambda x: x["week"])
                if weeks_list:
                    months.append({"month": oy_nom, "weeks": weeks_list})

            statistics = []
            total = 0
            for col in sorted(layout["stat_ustunlar"].keys()):
                info = layout["stat_ustunlar"][col]
                value = self._num_val(self._v(r, col))
                if isinstance(info, dict):
                    special = info.get("key")
                    if special == "total":
                        total = value
                        continue
                    statistics.append({
                        "key":   info.get("legendKey"),
                        "slug":  info.get("slug", ""),
                        "title": info.get("header", ""),
                        "value": value,
                    })
                else:
                    if info == "total":
                        total = value
                        continue
                    statistics.append({
                        "key":   None,
                        "slug":  str(info),
                        "title": "",
                        "value": value,
                    })

            weeks_dict = {str(k): v for k, v in sorted(weeks_xom.items())}

            courses.append({
                "course"    : kurs_val,
                "courseNum" : RIM_ARAB[kurs_val],
                "months"    : months,
                "weeks"     : weeks_dict,
                "total"     : total,
                "statistics": statistics,
            })

        return courses

    def _jami_o_qi(self):
        kurs_ust = self.layout.get("kurs_ustun", 1)

        for r in range(1, self.ws.max_row + 1):
            for c in range(kurs_ust, kurs_ust + 3):
                val = self._v(r, c)
                if val and _is_total_word(str(val).strip().lower()):
                    total = 0
                    statistics = []
                    for col in sorted(self.layout["stat_ustunlar"].keys()):
                        info = self.layout["stat_ustunlar"][col]
                        value = self._num_val(self._v(r, col))
                        if isinstance(info, dict):
                            if info.get("key") == "total":
                                total = value
                                continue
                            statistics.append({
                                "key":   info.get("legendKey"),
                                "slug":  info.get("slug", ""),
                                "title": info.get("header", ""),
                                "value": value,
                            })
                        else:
                            if info == "total":
                                total = value
                                continue
                            statistics.append({
                                "key":   None,
                                "slug":  str(info),
                                "title": "",
                                "value": value,
                            })
                    return {"total": total, "statistics": statistics}
        return {"total": 0, "statistics": []}

    SUMMARY_SHEET_KEYWORDS = ("xulosa", "summary", "резюме", "итог")

    def _xulosa_varaq_top(self):
        for nom in self.wb.sheetnames:
            low = nom.lower()
            if any(kw in low for kw in self.SUMMARY_SHEET_KEYWORDS):
                return self.wb[nom]
        return None

    def _xulosa_o_qi(self, legend_titles):
        ws = self._xulosa_varaq_top()
        if ws is None:
            return None

        natija = []
        for r in range(1, ws.max_row + 1):
            title_raw = ws.cell(r, 1).value
            weeks_raw = ws.cell(r, 2).value
            if title_raw is None or not isinstance(weeks_raw, (int, float)):
                continue
            title = str(title_raw).strip()
            if not title or _is_total_word(title.lower()):
                continue

            semester_raw = ws.cell(r, 3).value
            note_raw = ws.cell(r, 4).value
            natija.append({
                "key"     : _match_legend_key(title, legend_titles),
                "title"   : title,
                "weeks"   : self._num_val(weeks_raw),
                "semester": str(semester_raw).strip() if semester_raw is not None else None,
                "note"    : str(note_raw).strip() if note_raw not in (None, "") else None,
            })

        return natija if natija else None

    def parse(self):
        keys = self._legend_o_qi()
        legend_titles = {k["key"]: k["title"] for k in keys if k.get("title")}

        for col, info in self.layout["stat_ustunlar"].items():
            if not isinstance(info, dict):
                continue
            header = info.get("header", "")
            legend_key = _match_legend_key(header, legend_titles)
            if legend_key:
                info["legendKey"] = legend_key

        natija = {
            "metadata": {
                "file"        : Path(self.fayl_yoli).name,
                "coursesCount": 0,
                "weeksCount"  : len(self.layout["hafta_xaritasi"]),
                "monthsCount" : len(self.layout["oy_chegaralari"]),
            },
            "keys"      : keys,
            "courses"   : self._kurslar_o_qi(),
            "allValues" : self._jami_o_qi(),
            "summary"   : self._xulosa_o_qi(legend_titles),
            "warnings"  : self.warnings,
        }
        natija["metadata"]["coursesCount"] = len(natija["courses"])
        self.wb.close()
        return natija


if __name__ == "__main__":
    import argparse
    if sys.platform == "win32":
        try:
            sys.stdout.reconfigure(encoding="utf-8", errors="replace")
            sys.stderr.reconfigure(encoding="utf-8", errors="replace")
        except Exception:
            pass


    parser = argparse.ArgumentParser(description="Jarayon xlsx reader v3.5")
    parser.add_argument("fayl",              help="xlsx fayl yo'li")
    parser.add_argument("--sheet",           help="Sheet nomi (ixtiyoriy)", default=None)
    parser.add_argument("--stdout",          help="Natijani stdout ga chiqarish (JSON fayl emas)",
                        action="store_true")
    parser.add_argument("--out",             help="Chiqish JSON fayl nomi", default=None)
    parser.add_argument("--debug",           help="Layout debug ma'lumotini chiqarish",
                        action="store_true")
    args = parser.parse_args()

    reader = JarayonReader(args.fayl, sheet_nom=args.sheet)

    if args.debug and not args.stdout:
        l = reader.layout
        print("── Layout ──────────────────────────────────────", file=sys.stderr)
        print(f"  stat_ustunlar: {l['stat_ustunlar']}", file=sys.stderr)

    natija = reader.parse()

    if args.stdout:
        sys.stdout.write(json.dumps(natija, ensure_ascii=False))
        sys.stdout.flush()
    else:
        out = args.out or (Path(args.fayl).stem + "_v3_5.json")
        with open(out, "w", encoding="utf-8") as f:
            json.dump(natija, f, ensure_ascii=False, indent=2)

        m = natija["metadata"]
        print(f"✅ Saqlandi   : {out}")
        print(f"   Kurslar   : {m['coursesCount']}")
        print(f"   Haftalar  : {m['weeksCount']}")
        print(f"   Oylar     : {m['monthsCount']}")
        print(f"   Keys      : {len(natija['keys'])} belgi")
        print()
        print("── Keys ────────────────────────────────────────")
        for el in natija["keys"]:
            print(f"  [{el['key']}]  {el['title']}")
        print()
        print("── Courses ─────────────────────────────────────")
        for k in natija["courses"]:
            print(f"  {k['course']}-kurs: total={k['total']}, "
                  f"stats={list(k['statistics'].keys())}")
