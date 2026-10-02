#!/usr/bin/env python3
import sys, os, argparse, json
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from reja_pdf_reader import parse_pdf


def _thin_border():
    side = Side(style="thin", color="000000")
    return Border(left=side, right=side, top=side, bottom=side)


def _center(wrap=True):
    return Alignment(horizontal="center", vertical="center", wrap_text=wrap)


def _left(wrap=True):
    return Alignment(horizontal="left", vertical="center", wrap_text=wrap)


def _bold(size=10):
    return Font(bold=True, size=size)


def _font(size=10):
    return Font(size=size)


def _fill(color):
    return PatternFill(start_color=color, end_color=color, fill_type="solid")


def _set(ws, row, col, value, *, bold=False, size=10, align=None, bg=None, border=True):
    cell = ws.cell(row, col, value=value)
    cell.font = _bold(size) if bold else _font(size)
    cell.alignment = align or _center()
    if bg:
        cell.fill = _fill(bg)
    if border:
        cell.border = _thin_border()
    return cell


def _merge(ws, r1, c1, r2, c2, value=None, **style):
    ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
    _set(ws, r1, c1, value, **style)
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            if (r, c) != (r1, c1):
                ws.cell(r, c).border = _thin_border()


STD_MONTH_RANGES = [
    ("Sentabr", 1,  5),
    ("Oktabr",  6,  9),
    ("Noyabr",  10, 14),
    ("Dekabr",  15, 18),
    ("Yanvar",  19, 23),
    ("Fevral",  24, 27),
    ("Mart",    28, 32),
    ("Aprel",   33, 36),
    ("May",     37, 40),
    ("Iyun",    41, 44),
    ("Iyul",    45, 48),
    ("Avgust",  49, 52),
]

STAT_COLS = [
    (54, "Jami"),
    (55, "Nazariy va amaliy ta'lim"),
    (56, "Attestatsiyalar"),
    (57, "Kredit ta'lim tizimiga kirish"),
    (58, "Malakaviy amaliyot"),
    (59, "Yakuniy davlat attestatsiyasi"),
    (60, "Ta'til haftalari soni"),
    (61, "GPA ko'rsatkichini hisoblash"),
    (62, "Hammasi"),
]

LEGEND_POSITIONS = [
    (1,  2,  "",  "Nazariy va amaliy ta'lim"),
    (5,  6,  "A", "Attestatsiyalar"),
    (9,  10, "K", "Kredit ta'lim tizimiga kirish"),
    (15, 16, "M", "Malakaviy amaliyot"),
    (21, 22, "D", "Yakuniy Davlat attestatsiyasi"),
    (31, 32, "T", "Ta'til"),
    (37, 38, "G", "GPA ko'rsatkichini hisoblash"),
]


def _find_course_by_num(courses, num):
    for c in courses:
        if c.get("courseNum") == num:
            return c
    return None


def _find_stat_value(course, slug_or_key):
    stats = course.get("statistics", [])
    if isinstance(stats, list):
        for s in stats:
            if s.get("slug") == slug_or_key or s.get("key") == slug_or_key:
                return s.get("value", 0)
    elif isinstance(stats, dict):
        return stats.get(slug_or_key, 0)
    return 0


def generate_jarayoni_xlsx(parsed, output_path, num_courses=5):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Jadvali"

    _merge(ws, 1, 1, 1, 62, "I.  O'QUV JARAYONI JADVALI", bold=True, size=12, align=_center())

    _merge(ws, 2, 1, 6, 1, "Kurs", bold=True, size=10, align=_center())
    _merge(ws, 2, 2, 4, 53, "Haftalar", bold=True, size=10, align=_center())

    _merge(ws, 2, 54, 5, 54, "Jami", bold=True, size=8, align=_center())
    _merge(ws, 2, 55, 2, 59, "O'quv jarayoni, haftalari soni:", bold=True, size=8, align=_center())
    _merge(ws, 2, 60, 6, 60, "Ta'til haftalari soni", bold=True, size=7, align=_center())
    _merge(ws, 2, 61, 6, 61, "GPA ko'rsatkichini hisoblash", bold=True, size=7, align=_center())
    _merge(ws, 2, 62, 6, 62, "Hammasi", bold=True, size=8, align=_center())

    _merge(ws, 3, 55, 3, 59, "shundan", bold=True, size=8, align=_center())

    shundan_labels = [
        (55, "Nazariy va amaliy ta'lim"),
        (56, "Attestatsiyalar"),
        (57, "Kredit ta'lim tizimiga kirish"),
        (58, "Malakaviy amaliyot"),
        (59, "Yakuniy davlat attestatsiyasi"),
    ]
    for col, lbl in shundan_labels:
        _merge(ws, 4, col, 5, col, lbl, bold=True, size=7, align=_center())

    for name, s, e in STD_MONTH_RANGES:
        _merge(ws, 5, 1 + s, 5, 1 + e, name, bold=True, size=8, align=_center())

    for wk in range(1, 53):
        _set(ws, 6, 1 + wk, wk, size=7, align=_center())

    roman = {1: "I", 2: "II", 3: "III", 4: "IV", 5: "V"}
    parsed_courses = parsed.get("courses", [])

    for k_idx in range(1, num_courses + 1):
        r = 6 + k_idx
        _set(ws, r, 1, roman.get(k_idx, str(k_idx)), bold=True, size=9, align=_center())

        course = _find_course_by_num(parsed_courses, k_idx)
        weeks = course.get("weeks", {}) if course else {}
        for wk in range(1, 53):
            key = weeks.get(str(wk)) or ""
            _set(ws, r, 1 + wk, key.strip() if key and key != " " else "", size=8, align=_center())

        stat_mapping = [
            (54, "total"),
            (55, "theoreticalPractical"),
            (56, "certification"),
            (57, "creditSystem"),
            (58, "qualification"),
            (59, "final"),
            (60, "vacation"),
            (61, "gpa"),
            (62, "all"),
        ]
        if course:
            stats = course.get("statistics", [])
            sv = {}
            if isinstance(stats, list):
                for s in stats:
                    sl = s.get("slug") or ""
                    kk = s.get("key") or ""
                    v = s.get("value", 0)
                    if sl: sv[sl] = v
                    if kk: sv[kk] = v

            col_slugs = {
                54: ["total", "jami"],
                55: ["theoreticalPractical", "theoreticalScientific",
                     "nazariy_va_amaliy_talim", "nazariy_talim_va_ilmiy_faoliyat", " "],
                56: ["certification", "attestatsiyalar", "A"],
                57: ["creditSystem", "kredit_talim_tizimiga_kirish", "K"],
                58: ["qualification", "qualificationPractice", "malakaviy_amaliyot", "M"],
                59: ["final", "finalStateAttestation", "yakuniy_davlat_attestatsiyasi", "D"],
                60: ["vacation", "tatil_haftalari_soni", "tatillar", "T"],
                61: ["gpa", "gpaCalculation", "gpa_korsatkichini_hisoblash", "G"],
                62: ["all", "hammasi"],
            }
            for col, slugs in col_slugs.items():
                val = 0
                for sl in slugs:
                    if sl in sv:
                        val = sv[sl]
                        break
                if col == 54 and not val:
                    val = course.get("total", 0)
                _set(ws, r, col, val if val else "", size=8, align=_center())

    _set(ws, 12, 1, "Jami", bold=True, size=9, align=_center())
    av = parsed.get("allValues", {})
    av_stats = av.get("statistics", []) if isinstance(av, dict) else []
    av_sv = {}
    if isinstance(av_stats, list):
        for s in av_stats:
            sl = s.get("slug") or ""
            kk = s.get("key") or ""
            v = s.get("value", 0)
            if sl: av_sv[sl] = v
            if kk: av_sv[kk] = v

    total_row_mapping = {
        54: ["total", "jami"],
        55: ["theoreticalScientific", "theoreticalPractical",
             "nazariy_va_amaliy_talim", " "],
        56: ["certification", "attestatsiyalar", "A"],
        57: ["creditSystem", "kredit_talim_tizimiga_kirish", "K"],
        58: ["qualification", "malakaviy_amaliyot", "M"],
        59: ["final", "finalStateAttestation", "yakuniy_davlat_attestatsiyasi", "D"],
        60: ["vacation", "tatil_haftalari_soni", "T"],
        61: ["gpa", "gpa_korsatkichini_hisoblash", "G"],
        62: ["all", "hammasi"],
    }
    _merge(ws, 12, 2, 12, 53, "", bold=True)
    for col, slugs in total_row_mapping.items():
        val = 0
        for sl in slugs:
            if sl in av_sv:
                val = av_sv[sl]
                break
        if col == 54 and not val:
            val = av.get("total", 0)
        _set(ws, 12, col, val if val else "", bold=True, size=9, align=_center())

    for c in range(1, 63):
        _set(ws, 13, c, "", border=False)

    parsed_keys = parsed.get("keys", [])
    key_title_map = {k.get("key", ""): k.get("title", "") for k in parsed_keys}

    legend_entries = [
        (2, 3, "",  key_title_map.get(" ", "Nazariy va amaliy ta'lim")),
        (6, 7, "A", key_title_map.get("A", "Attestatsiyalar")),
        (10, 13, "K", key_title_map.get("K", "Kredit ta'lim tizimiga kirish")),
        (16, 19, "M", key_title_map.get("M", "Malakaviy amaliyot")),
        (22, 26, "D", key_title_map.get("D", "Yakuniy Davlat attestatsiyasi")),
        (29, 30, "T", key_title_map.get("T", "Ta'til")),
        (34, 38, "G", key_title_map.get("G", "GPA ko'rsatkichini hisoblash")),
    ]
    for text_start, text_end, letter, title in legend_entries:
        sq_col = text_start - 1
        _set(ws, 14, sq_col, letter, bold=True, size=9, align=_center())
        _merge(ws, 14, text_start, 14, text_end, title, size=8, align=_center())

    ws.column_dimensions["A"].width = 6
    for c in range(2, 54):
        ws.column_dimensions[get_column_letter(c)].width = 2.5
    for c in range(54, 63):
        ws.column_dimensions[get_column_letter(c)].width = 7

    ws.row_dimensions[1].height = 20
    ws.row_dimensions[5].height = 20
    for r in range(2, 7):
        ws.row_dimensions[r].height = 15
    for r in range(7, 13):
        ws.row_dimensions[r].height = 18
    ws.row_dimensions[14].height = 25

    wb.save(output_path)
    return output_path


STANDARD_PARTICLE_COLUMNS = [
    ("soat",         "soat"),
    ("foiz",         "%"),
    ("jami",         "Jami"),
    ("maruza",       "Ma'ruza"),
    ("amaliy",       "Amaliy mashg' ulot"),
    ("laboratoriya", "Laboratoriya mashg' uloti"),
    ("seminar",      "Seminar"),
    ("mustaqil",     "Mustaqil ta' lim"),
]


def _particle_value(particle_list, slug):
    if isinstance(particle_list, list):
        for p in particle_list:
            if p.get("slug") == slug:
                return p.get("value", 0)
    elif isinstance(particle_list, dict):
        legacy_map = {
            "soat": "hour", "foiz": "percent", "jami": "total",
            "maruza": "lecture", "amaliy": "practical",
            "laboratoriya": "laboratory", "seminar": "seminar",
            "mustaqil": "independent",
        }
        return particle_list.get(legacy_map.get(slug, slug), 0)
    return 0


def generate_reja_xlsx(parsed, output_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "O'quv Rejasi"

    meta = parsed.get("meta", {})
    blocks = parsed.get("blocks", [])

    particle_items = [{"slug": s, "title": t} for s, t in STANDARD_PARTICLE_COLUMNS]
    meta_items = meta.get("particles", {}).get("items", []) or []
    meta_by_slug = {it.get("slug"): it.get("title") for it in meta_items if it.get("slug")}
    for it in particle_items:
        if it["slug"] in meta_by_slug and meta_by_slug[it["slug"]]:
            it["title"] = meta_by_slug[it["slug"]]

    max_sem = 10
    for b in blocks:
        sems = b.get("semesters", {})
        if isinstance(sems, dict):
            for k in sems.keys():
                try:
                    max_sem = max(max_sem, int(k))
                except (ValueError, TypeError):
                    pass
        for s in b.get("sciences", []):
            for k in s.get("semesters", {}).keys():
                try:
                    max_sem = max(max_sem, int(k))
                except (ValueError, TypeError):
                    pass

    kurs_count = max(1, (max_sem + 1) // 2)

    P_COUNT = max(len(particle_items), 8)
    P_START = 4
    P_END = P_START + P_COUNT - 1

    SOAT_START = P_END + 1
    SOAT_END = SOAT_START + max_sem - 1
    KRED_START = SOAT_END + 1
    KRED_END = KRED_START + max_sem - 1
    JK_COL = KRED_END + 1

    total_cols = JK_COL

    _merge(ws, 1, 1, 1, total_cols, "II. O'QUV REJASI", bold=True, size=14, align=_center())

    _merge(ws, 2, 1, 8, 1, "T/r", bold=True, size=10, align=_center())
    _merge(ws, 2, 2, 8, 2, "Fanning malakaviy kodi", bold=True, size=9, align=_center())
    _merge(ws, 2, 3, 8, 3, meta.get("title") or "O'quv bloklari, fanlar va faoliyat turlarining nomlari",
           bold=True, size=9, align=_center())

    p_title = meta.get("particles", {}).get("title") or "Talabaning o'quv yuklamasi (soatlarda)"
    _merge(ws, 2, P_START, 2, P_END, p_title, bold=True, size=9, align=_center())

    if P_COUNT >= 2:
        _merge(ws, 3, P_START, 6, P_START + 1, "Umumiy\nyuklamaning\nhajmi",
               bold=True, size=8, align=_center())
    if P_COUNT >= 7:
        _merge(ws, 3, P_START + 2, 3, P_START + 6, "Auditoriya mashg'ulotlari, soatlarda",
               bold=True, size=8, align=_center())
    if P_COUNT >= 8:
        _merge(ws, 3, P_START + 7, 8, P_START + 7, particle_items[7].get("title") if len(particle_items) > 7 else "Mustaqil ta' lim",
               bold=True, size=8, align=_center())

    for i, it in enumerate(particle_items[:P_COUNT]):
        col = P_START + i
        title = it.get("title") or it.get("slug", "")
        if i < 2:
            continue
        if it["slug"] == "mustaqil":
            continue
        _merge(ws, 4, col, 8, col, title, bold=True, size=7, align=_center())

    if P_COUNT >= 1:
        _set(ws, 7, P_START, particle_items[0].get("title") or "soat",
             bold=True, size=8, align=_center())
    if P_COUNT >= 2:
        _set(ws, 7, P_START + 1, particle_items[1].get("title") or "%",
             bold=True, size=8, align=_center())

    _merge(ws, 2, SOAT_START, 2, SOAT_END, "Soatlarning kurslar, semestrlar va haftalar bo'yicha taqsimoti",
           bold=True, size=9, align=_center())

    for k in range(kurs_count):
        c1 = SOAT_START + k * 2
        c2 = c1 + 1
        if c2 > SOAT_END:
            c2 = SOAT_END
        _merge(ws, 3, c1, 3, c2, f"{k+1}-kurs", bold=True, size=9, align=_center())

    for k in range(kurs_count):
        c1 = SOAT_START + k * 2
        c2 = min(c1 + 1, SOAT_END)
        _merge(ws, 4, c1, 4, c2, "30", size=9, align=_center())

    _merge(ws, 5, SOAT_START, 5, SOAT_END, "Semestrlar", bold=True, size=9, align=_center())

    for s in range(max_sem):
        _set(ws, 6, SOAT_START + s, s + 1, size=9, align=_center())

    _merge(ws, 7, SOAT_START, 7, SOAT_END,
           "Semestrdagi auditoriya mashgʻulotlari haftalarining  soni",
           bold=True, size=7, align=_center())

    for s in range(max_sem):
        _set(ws, 8, SOAT_START + s, "15", size=9, align=_center())

    _merge(ws, 2, KRED_START, 2, KRED_END,
           "Kreditlarning kurslar, semestrlar va haftalar bo'yicha taqsimoti",
           bold=True, size=9, align=_center())

    for k in range(kurs_count):
        c1 = KRED_START + k * 2
        c2 = min(c1 + 1, KRED_END)
        _merge(ws, 3, c1, 3, c2, f"{k+1}-kurs", bold=True, size=9, align=_center())
        _merge(ws, 4, c1, 4, c2, "30", size=9, align=_center())

    _merge(ws, 5, KRED_START, 5, KRED_END, "Semestrlar", bold=True, size=9, align=_center())
    for s in range(max_sem):
        _set(ws, 6, KRED_START + s, s + 1, size=9, align=_center())

    _merge(ws, 7, KRED_START, 7, KRED_END, "Kredit taqsimoti", bold=True, size=8, align=_center())
    for s in range(max_sem):
        _set(ws, 8, KRED_START + s, "30", size=9, align=_center())

    _merge(ws, 2, JK_COL, 7, JK_COL, "Jami kreditlar", bold=True, size=9, align=_center())
    _set(ws, 8, JK_COL, "300", size=9, align=_center())

    for c in range(1, total_cols + 1):
        _set(ws, 9, c, c, size=8, align=_center())

    cur_row = 10
    for block in blocks:
        _set(ws, cur_row, 1, block.get("serialNumber", ""),
             bold=True, size=9, align=_center(), bg="E8F0F8")
        _set(ws, cur_row, 2, block.get("blockCode", "") or block.get("code", ""),
             bold=True, size=9, align=_center(), bg="E8F0F8")
        _set(ws, cur_row, 3, block.get("title", ""),
             bold=True, size=9, align=_left(), bg="E8F0F8")

        b_particle = block.get("particle", [])
        for i, it in enumerate(particle_items[:P_COUNT]):
            val = _particle_value(b_particle, it["slug"])
            _set(ws, cur_row, P_START + i, val if val else "",
                 bold=True, size=8, align=_center(), bg="E8F0F8")

        b_sems = block.get("semesters", {})
        for s in range(max_sem):
            sd = b_sems.get(str(s + 1), {}) if isinstance(b_sems, dict) else {}
            hour_val = sd.get("hour", 0) if isinstance(sd, dict) else 0
            cred_val = sd.get("credit", 0) if isinstance(sd, dict) else 0
            _set(ws, cur_row, SOAT_START + s, hour_val if hour_val else "",
                 bold=True, size=8, align=_center(), bg="E8F0F8")
            _set(ws, cur_row, KRED_START + s, cred_val if cred_val else "",
                 bold=True, size=8, align=_center(), bg="E8F0F8")

        _set(ws, cur_row, JK_COL, block.get("totalCredit", 0),
             bold=True, size=9, align=_center(), bg="E8F0F8")

        cur_row += 1

        for sci in block.get("sciences", []):
            _set(ws, cur_row, 1, sci.get("serialNumber", ""), size=8, align=_center())
            _set(ws, cur_row, 2, sci.get("code", ""), size=8, align=_center())
            _set(ws, cur_row, 3, sci.get("title", ""), size=8, align=_left())

            s_particle = sci.get("particle", [])
            for i, it in enumerate(particle_items[:P_COUNT]):
                val = _particle_value(s_particle, it["slug"])
                _set(ws, cur_row, P_START + i, val if val else "",
                     size=8, align=_center())

            s_sems = sci.get("semesters", {})
            for s in range(max_sem):
                sd = s_sems.get(str(s + 1), {}) if isinstance(s_sems, dict) else {}
                hour_val = sd.get("hour", 0) if isinstance(sd, dict) else 0
                cred_val = sd.get("credit", 0) if isinstance(sd, dict) else 0
                _set(ws, cur_row, SOAT_START + s, hour_val if hour_val else "",
                     size=8, align=_center())
                _set(ws, cur_row, KRED_START + s, cred_val if cred_val else "",
                     size=8, align=_center())

            _set(ws, cur_row, JK_COL, sci.get("totalCredit", 0), size=8, align=_center())
            cur_row += 1

    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 35
    for c in range(P_START, P_END + 1):
        ws.column_dimensions[get_column_letter(c)].width = 9
    for c in range(SOAT_START, total_cols + 1):
        ws.column_dimensions[get_column_letter(c)].width = 5

    ws.row_dimensions[1].height = 22
    for r in range(2, 9):
        ws.row_dimensions[r].height = 18
    ws.row_dimensions[3].height = 35
    for r in range(10, cur_row):
        ws.row_dimensions[r].height = 18

    wb.save(output_path)
    return output_path


def main():
    if sys.platform == "win32":
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")

    ap = argparse.ArgumentParser(description="PDF reja → 2 ta XLSX (jarayoni + rejasi)")
    ap.add_argument("pdf_path", help="PDF fayl yo'li")
    ap.add_argument("--out-dir", default=".", help="Chiqish katalogi (default: .)")
    ap.add_argument("--prefix", default=None, help="Fayl nomi prefiksi (default: PDF nomidan)")
    args = ap.parse_args()

    if not os.path.exists(args.pdf_path):
        print(f"Fayl topilmadi: {args.pdf_path}", file=sys.stderr)
        sys.exit(1)

    print(f"PDF o'qilmoqda: {args.pdf_path}", file=sys.stderr)
    parsed = parse_pdf(args.pdf_path)
    print(f"  bloklar: {len(parsed.get('blocks', []))}", file=sys.stderr)
    print(f"  kurslar: {len(parsed.get('courses', []))}", file=sys.stderr)
    print(f"  keys: {len(parsed.get('keys', []))}", file=sys.stderr)

    base = args.prefix or Path(args.pdf_path).stem
    out_dir = Path(args.out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    j_out = out_dir / f"{base}_jarayoni.xlsx"
    r_out = out_dir / f"{base}_rejasi.xlsx"

    print(f"\nJarayoni XLSX yaratilmoqda: {j_out}", file=sys.stderr)
    generate_jarayoni_xlsx(parsed, str(j_out))
    print(f"  ✓ saqlandi ({j_out.stat().st_size // 1024} KB)", file=sys.stderr)

    print(f"\nRejasi XLSX yaratilmoqda: {r_out}", file=sys.stderr)
    generate_reja_xlsx(parsed, str(r_out))
    print(f"  ✓ saqlandi ({r_out.stat().st_size // 1024} KB)", file=sys.stderr)

    print(f"\n2 ta XLSX tayyor:", file=sys.stderr)
    print(f"  {j_out}", file=sys.stderr)
    print(f"  {r_out}", file=sys.stderr)


if __name__ == "__main__":
    main()
